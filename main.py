#!/usr/bin/env python3

from docxtpl import DocxTemplate
import datetime
import os
import google.generativeai as genai
from google.api_core.exceptions import ResourceExhausted
import pyperclip
import subprocess
from docx import Document
from tqdm import tqdm
from selenium import webdriver
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.common.keys import Keys
from selenium.common.exceptions import NoSuchElementException
from selenium.webdriver.common.by import By
import time
import json
import re
import sys
from urllib.parse import urlparse

api_keys = [""]
active_key_index = 0 

def convert_word_to_pdf(input_file, output_dir):
    LIBREOFFICE_BINARY = '/Applications/LibreOffice.app/Contents/MacOS/soffice'
    subprocess.run(
        [LIBREOFFICE_BINARY, '--headless', '--convert-to', 'pdf', '--outdir', output_dir, input_file],
        stdout=subprocess.PIPE, stderr=subprocess.PIPE
    )

def delete_file(file_path):
    try:
        os.remove(file_path)
    except OSError as e:
        print(f"Error deleting file {file_path}: {e.strerror}")

def extract_text_from_docx(docx_file):
    doc = Document(docx_file)
    full_text = [para.text for para in doc.paragraphs]
    return '\n'.join(full_text).strip()

def extract_json(response):
    """
    Extract JSON content from a string, stripping everything except for the content
    inside the curly brackets.
    """
    json_match = re.search(r'\{.*\}', response, re.DOTALL)
    if json_match:
        return json_match.group(0)  # Return only the matched JSON content
    else:
        raise ValueError("No valid JSON found in response")

def is_valid_url(url):
    parsed_url = urlparse(url)
    return all([parsed_url.scheme, parsed_url.netloc])

def short_form_position_name(position_name):
    # Split the position name into words
    words = position_name.split()

    # Initialize an empty string to hold the short form
    short_form = ""

    # List of words to skip (intern, internship, co-op, etc.)
    skip_words = ["intern", "internship", "co-op", "coop", "student"]

    # Loop through each word in the position name
    for word in words:
        # Skip words that contain numbers or special characters
        if not word.isalpha():
            continue

        # Check for specific words and apply custom logic
        if word.lower() == "software":
            # If the word is "Software", append "SW" to the short form
            short_form += "SW"
        elif word.lower() in skip_words:
            # Skip the word if it's in the skip_words list
            continue
        else:
            # Otherwise, append the first letter of the word
            short_form += word[0].upper()

    return short_form

def generate_with_gemini(prompt, retries=3):
    """
    Generate content using the Gemini API, using the last successful API key until it gets exhausted.
    """
    global active_key_index  # Use the global variable to track the current API key
    
    # Configure the model with the current active API key
    genai.configure(api_key=api_keys[active_key_index])
    model = genai.GenerativeModel("gemini-1.5-flash")
    
    for attempt in range(retries):
        try:
            # Try generating content using the active API key
            response = model.generate_content(prompt)
            return response.text  # Successfully generated content

        except ResourceExhausted as e:
            print(f"Quota exhausted for API key {active_key_index + 1}.")
            
            # Switch to the other API key
            active_key_index = 1 if active_key_index == 0 else 0  # Toggle between 0 and 1
            
            # Reconfigure the model with the new API key
            print(f"Switching to API key {active_key_index + 1}...")
            genai.configure(api_key=api_keys[active_key_index])
            model = genai.GenerativeModel("gemini-1.5-flash")

            # Continue retrying with the new key
        except Exception as e:
            print(f"An error occurred: {e}")
            raise e

    # If we get here after retries, it means the requ est failed after all retries
    print("Max retries reached. Could not complete the request.")
    raise Exception("Max retries reached.")

def get_job_details(url, pbar, max_retries=3):
    for attempt in range(1, max_retries + 1):
        # Step 1: Set up Selenium with headless Chrome
        pbar.set_description(f"Initializing WebDriver")
        chrome_options = Options()
        chrome_options.add_argument("--headless")
        chrome_options.add_argument("--disable-gpu")
        chrome_options.add_argument("--incognito")
        driver = webdriver.Chrome(options=chrome_options)
        if attempt == 1:
            pbar.update(1)

        # Check if the URL is from LinkedIn
        
        parsed_url = urlparse(url)
        if "linkedin.com" in parsed_url.netloc:
            pbar.set_description(f"Navigating to the LinkedIn posting (Attempt {attempt})")
            driver.get(url)

            # Wait for the page to load
            time.sleep(1)

            if attempt == 1:
                pbar.update(1)

            # Step 2: Check if the current URL matches the intended URL (to check for redirection to login)
            current_url = driver.current_url
            if not current_url.startswith(url):
                pbar.set_description(f"Page redirected, reloading (Attempt {attempt})")

                # Attempt to reload the page without clearing history
                driver.get(url)
                time.sleep(1)
                current_url = driver.current_url

                # If still redirected, continue to the next attempt (quit and retry)
                if not current_url.startswith(url):
                    pbar.set_description(f"Still redirected, quitting WebDriver (Attempt {attempt})")
                    driver.quit()
                    if attempt == max_retries:
                        # If max retries reached, skip this URL
                        pbar.set_description(f"Max retries reached. Skipping this URL.")
                        return None, None, None, None
                    continue

            # If the URL is correct, proceed to extract job details
            if current_url.startswith(url):
                pbar.set_description(f"Extracting job details")

                try:
                    job_description_element = WebDriverWait(driver, 10).until(
                        EC.presence_of_element_located((By.TAG_NAME, 'body'))
                    )

                    # Get the text content from the job description and remove extra whitespace
                    job_description_text = job_description_element.text.strip()

                    # Extract all <ul> and <li> elements inside this job description section
                    ul_elements = job_description_element.find_elements(By.TAG_NAME, 'ul')  # Find all unordered lists
                    list_text = ""

                    # Iterate over each <ul> and extract the text from <li> elements
                    for ul in ul_elements:
                        li_elements = ul.find_elements(By.TAG_NAME, 'li')  # Find all list items inside the <ul>
                        for li in li_elements:
                            list_text += li.get_attribute("innerText").strip() + "\n"  # Add each list item to the string, trimming whitespace

                    # Combine the text and list items into one string, removing excessive line breaks and extra spaces
                    job_description_text_cleaned = "\n".join([line.strip() for line in job_description_text.splitlines() if line.strip()])

                    # Clean up the list text to remove excessive whitespace or empty lines
                    list_text_cleaned = "\n".join([line.strip() for line in list_text.splitlines() if line.strip()])

                    # Combine the cleaned job description and list items
                    job_description = job_description_text_cleaned + "\n\nImportant Items, could potentially be technical skills:\n" + list_text_cleaned

                    # Truncate the job description to 6000 characters if necessary
                    job_description = job_description[:6000]

                except Exception as e:
                    print(f"Error extracting job details: {e}")
                    job_description = None

            else:
                # If retries limit is reached and still redirected, return None
                return None, None, None, None

        else:
            # For non-LinkedIn URLs, just extract the body content
            pbar.set_description(f"Navigating to the job URL")
            driver.get(url)

            time.sleep(1)

            if attempt == 1:
                pbar.update(1)

            try:
                body_element = WebDriverWait(driver, 10).until(
                    EC.presence_of_element_located((By.TAG_NAME, "body"))
                )
                body_text = body_element.text
                job_description = body_text[:6000] 

            except Exception as e:
                print(f"Error extracting body content: {e}")
                job_description = None

        # Use AI to deduce the company name and position
        try:
            ai_prompt = f"""
Extract the company name and position title from the following text.

Text:
{job_description}

Instructions:
- Focus **only** on the core job posting content. Ignore sections related to government forms, surveys, legal disclaimers, and any non-job-related information.
- The **company name** and **job title** are typically mentioned at the beginning of the job posting. Extract these from the relevant section of the text.
- The **company name** will typically be a name you recognize, it won't be a generic term like "Engineering" or "Software", it would specifically be a name of a company.
  - Note that if the company name repeatedly says "LinkedIn" or "Simplify", this may be because that is the site that is hosting the posting and not the company itself.
  - The company is not typically named LinkedIn, do not name it LinkedIn unless you are absolutely confident
- For the **company name**, simplify it by removing legal terms like 'Inc.', 'LLC', 'Ltd.', or words like 'Markets', 'Corporation', etc.
- For the **position title**, rephrase it into a standard format. Use terms like 'Software Engineering Intern' or 'Software Engineer', removing extra details like hyphens, department, location, or year.
  - Ensure that the position title is phrased so it fits naturally in the sentence, “I'm excited to join the team as a {{position}}.”
- Also, for the **position title**, if the job posting is **overly specific or niche** (e.g., "Front-End Software Engineer District 2 Webview Editing"), simplify it to the **general role** (e.g., "Front-End Software Engineer"). Focus on the broader job role instead of hyper-specific department names, project names, or locations.
  - For example, "Software Engineering Intern, Web" should be rephrased to "Software Engineering Intern"
  - For example, "Software Developer Intern - Cybersecurity" should be rephrased to "Software Developer Intern"
  - For example, "2025 Summer Software Engineer" should be rephrased to "Software Engineer"
- For the **position title**, ensure that the word "co-op" is replaced with the word "intern"
  - For example, "Software Developer Co-op" should be rephrased to "Software Developer Intern"
- For the **requirements**, try to find technical skills in the job description. If you can't find any, think of specific requirements for the type of job
- You should **ignore** sections that contain unrelated information such as:
  - Legal disclaimers
  - Equal Employment Opportunity (EEO) statements
  - Government forms, surveys, or voluntary self-identification forms
  - Privacy policy notices or links
  - Any additional instructions or forms

Provide the output in the **exact JSON format** below without any extra explanations:
{{
    "company_name": "Simplified Company Name",
    "position_name": "Simplified Position Title",
    "requirements": "List of Requirements"
}}

Ensure the response is valid JSON.
"""
            if attempt == 1:
                pbar.set_description(f"Extracting job details")
            else:
                pbar.set_description(f"Extracting job details (attempt {attempt})")

            ai_content = generate_with_gemini(ai_prompt.strip())

            # Try to extract JSON from AI response
            json_start = ai_content.index('{')
            json_end = ai_content.rindex('}') + 1
            json_str = ai_content[json_start:json_end]
            ai_data = json.loads(json_str)
            company_name = ai_data.get('company_name', '').strip()
            position_name = ai_data.get('position_name', '').strip()
            requirements = ai_data.get('requirements', '').strip()

            if company_name and position_name:
                # Successful extraction
                driver.quit()
                pbar.set_description("Successfully extracted job details")
                pbar.update(3)
                return company_name, position_name, requirements, job_description
            else:
                pbar.set_description(f"Failed to extract details, retrying")
                driver.quit()
                if attempt < max_retries:
                    time.sleep(1)  # Optional delay between retries
                else:
                    pbar.set_description("Max retries reached. Skipping this URL")
                    return None, None, None, None

        except Exception as e:
            pbar.set_description(f"Error during job detail extraction: {e}")
            driver.quit()
            if attempt < max_retries:
                pbar.set_description("Retrying...")
                time.sleep(1)  # Optional delay between retries
            else:
                pbar.set_description("Max retries reached. Skipping this URL")
                return None, None, None, None

def main():

    first_run = True

    while True:
        # Get the directory where the script is located
        script_dir = os.path.dirname(os.path.realpath(__file__))

        # Construct the path to the Word template file
        template_path = os.path.join(script_dir, "Template.docx")

        url = ""
        placeholder = ""
        company_name, position_name, requirements, job_description = None, None, None, None

        while True:
            # If not regenerating, ask for a new URL or input
            if first_run:
                # Force valid URL entry, ignore 'r' on first run
                url = input("Enter the job posting URL: ").strip()
                while not is_valid_url(url):
                    if url.lower() == 'r':
                        print("You cannot regenerate on the first run. Please enter a valid job posting URL.")
                    else:
                        print("Invalid URL. Please enter a valid URL.")
                    url = input("Enter the job posting URL: ").strip()
            else:
                placeholder = input("Enter the job posting URL or type 'r' to regenerate: ").strip()
                while placeholder.lower() != 'r' and not is_valid_url(placeholder):
                    print("Invalid entry. Please enter a valid URL.")
                    placeholder = input("Enter the job posting URL or type 'r' to regenerate: ").strip()
                if placeholder != "r":
                    url = placeholder
            
            first_run = False  # Set to False after the first run

            total_steps = 15  # Adjusted total steps for progress bar
            with tqdm(total=total_steps) as pbar:
                pbar.set_description("Starting")

                company_name, position_name, requirements, job_description = get_job_details(url, pbar)

                manual_entry = False  # Flag to check if user entered details manually

                # Check for manual entry if extraction failed
                if not company_name or not position_name:
                    print("Error: Unable to extract company name or position title.")
                    user_choice = input("Do you want to enter the company name and position title manually? (y/n): ").strip().lower()
                    if user_choice == 'y':
                        company_name = input("Enter Company Name: ").strip()
                        position_name = input("Enter Position Title: ").strip()
                        manual_entry = True
                    else:
                        print("Exiting script.")
                        sys.exit(1)

                if not job_description or not job_description.strip():
                    print("Error: Unable to extract job description.")
                    user_choice = input("Do you want to paste the job description manually? (y/n): ").strip().lower()
                    if user_choice == 'y':
                        print("Please paste the job description below. Press Enter twice when done:")
                        lines = []
                        while True:
                            line = input()
                            if line:
                                lines.append(line)
                            else:
                                break
                        job_description = '\n'.join(lines)
                    else:
                        print("Exiting script.")
                        sys.exit(1)

                if manual_entry:
                    print(f"Company Name: {company_name}")
                    print(f"Position Name: {position_name}")
                pbar.update(1)

                today_date = datetime.datetime.today().strftime("%B %d, %Y")
                fname_date = datetime.datetime.today().strftime("%Y-%m-%d")
                pbar.update(1)

                short_form = short_form_position_name(position_name)

                extension = "'s" if not company_name.endswith('s') else "’"
                company_name_plural = company_name + extension

                # AI generation for responseTop
                pbar.set_description(f"Generating {short_form} CL")
                responseTop_prompt = f"""
Using the job description provided below, generate a final concluding sentence for a paragraph in my cover letter. This sentence should highlight how my skills and qualifications align with the job description, technical skills, and requirements.

**Important guidelines**:
- **Do not copy** or repeat exact phrases from the job description, technical skills, or requirements.
- Extract relevant technical and soft skills from: "{requirements}".
    - If no skills are found, move on without adding any.
    - **Use specific skills or technologies mentioned in the job description**, such as programming languages, frameworks, or tools, instead of using placeholders like "[Programming Language]" or "[Skills]".
    - No square brackets or generic placeholders should appear in the final sentence.
- **Summarize** how my background fits the role without repeating the job description's details.
- **Do not describe the company** or the position.
- Refer to the company as "{company_name}".
    - Use "{company_name_plural}" as the plural.
- The sentence **must** start with **"I am eager to leverage..."**.
Your sentence should **add value** to the existing content of the cover letter and conclude the paragraph meaningfully.

**Important**: The sentence must be complete, with no need for further input.
**Important**: You must **only** return the result in **valid JSON format**. No additional text, explanations, or commentary should be included—**just the JSON response**. The response must follow this structure exactly, including the correct curly braces and the key **"responseTop"**.

Job Description:
{job_description}

If the output is not in valid JSON format, the response will be considered incorrect.

Return the response **only** in the following format:
```json
{{
  "responseTop": "Generated sentence here."
}}
"""
                responseTop = generate_with_gemini(responseTop_prompt.strip())

                 # Extract and parse the JSON content
                try:
                    json_content = extract_json(responseTop.strip())
                    response_top_json = json.loads(json_content)
                    response_top_sentence = response_top_json['responseTop']
                except ValueError as e:
                    print(f"Error parsing JSON for responseTop: {e}")
                    return
                
                pbar.update(4)
                pbar.set_description(f"Glazing {company_name}")

                # AI generation for glazing
                glazing_prompt = f"""
Using the company values and goals provided in the job description below, generate a paragraph for my cover letter that highlights how my personal values and professional goals align with the company's motives and objectives.

**Important guidelines**:
- **Do not copy or repeat** the company's values or goals word-for-word from the job description, if necessary: rephrase.
- **Paraphrase and rephrase** the company's values/goals, and focus on aligning them with my own values and ambitions.
- **Do not describe the company** in detail (e.g., no mentioning of their history, products, or services).
- Refer to the company as "{company_name}".
    - Use "{company_name_plural}" as the plural.
- **Limit the response to 2-3 sentences only**. Keep the paragraph concise and to the point.
- The sentence **must** start with **"I am drawn by...."**.
Your paragraph should demonstrate how my personal goals, values, and professional mission align with the company's broader objectives and should flow smoothly as part of a professional cover letter.

**Important**: You must **only** return the result in **valid JSON format**. No other text, explanations, or commentary should be included—**just the JSON response**. The response must be structured exactly like the format below, including the correct curly braces and the key **"glazing"**.

Job Description:
{job_description}

If the output is not in valid JSON format, the response will be considered incorrect.
Avoid overlap in content, structure, or concepts with "{response_top_sentence}", ensuring no similarities in phrasing, wording or ideas.

Return the response **only** in the following format:
```json
{{
  "glazing": "Generated paragraph here."
}}
"""
                glazing = generate_with_gemini(glazing_prompt.strip())

                # Extract and parse the JSON content
                try:
                    json_content = extract_json(glazing.strip())
                    glazing_json = json.loads(json_content)
                    glazing_paragraph = glazing_json['glazing']
                except ValueError as e:
                    print(f"Error parsing JSON for glazing: {e}")
                    return

                pbar.update(1)

                a = "an" if position_name and position_name[0].upper() in ["A", "E", "I", "O", "U"] else "a"

                context = {
                    "today_date": today_date,
                    "position_name": position_name,
                    "company_name": company_name,
                    "company_name_plural": company_name_plural,
                    "a": a,
                    "generate": response_top_sentence,
                    "glazing": glazing_paragraph,
                }
                pbar.update(1)

                # Opening template
                pbar.set_description("Creating Word document")
                doc = DocxTemplate(template_path)

                # Load document
                doc.render(context)

                output_file_path = os.path.join(script_dir, f"{company_name} {position_name} {fname_date}.docx")
                doc.save(output_file_path)
                pbar.update(1)

                # Extract text from the updated DOCX file and copy it to the clipboard
                document_text = extract_text_from_docx(output_file_path)
                pyperclip.copy(document_text)

                # Convert DOCX to PDF using LibreOffice
                pbar.set_description("Converting to PDF")
                convert_word_to_pdf(output_file_path, script_dir)

                # Delete the original DOCX file
                delete_file(output_file_path)
                pbar.update(1)

                pbar.set_description("Done")

if __name__ == "__main__":
    main()